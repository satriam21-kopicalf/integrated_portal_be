"""Excel export job tests (database stubbed, jobs run inline, see conftest.py)."""
import io
import json
import os
from datetime import datetime, timedelta, timezone

import openpyxl

from app import exports, gsheets
from app.config import get_settings
from app.esb_report import REPORT_HEADERS


def _start(client, **body):
    res = client.post("/api/exports", json=body)
    assert res.status_code == 202, res.text
    return client.get(f"/api/exports/{res.json()['id']}").json()


def _workbook(client, job):
    res = client.get(job["downloadUrl"])
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert job["fileName"] in res.headers["content-disposition"]
    assert "content-encoding" not in res.headers  # not gzipped
    return openpyxl.load_workbook(io.BytesIO(res.content))


def _data_rows(ws, header_row):
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    assert list(rows[0]) == REPORT_HEADERS
    return rows[1:]


def test_export_matches_esb_layout(client):
    job = _start(client, dateFrom="2026-09-29", dateTo="2026-09-30")
    assert job["status"] == "done" and job["type"] == "sales"
    assert (job["totalDays"], job["daysDone"], job["headers"], job["rows"], job["sheets"]) == (2, 2, 2, 4, 1)
    assert job["fileName"] == "Sales_Recapitulation_Detail_2026-09-29_to_2026-09-30.xlsx"

    wb = _workbook(client, job)
    assert wb.sheetnames == ["Report", "Ringkasan"]
    ws = wb["Report"]
    assert ws["A1"].value == "Sales Recapitulation Detail Report" and ws["A2"].value == "PT Yuda Prawira Group"
    assert ws["A5"].value == "Period" and ws["B5"].value == "29-09-2026 - 30-09-2026"
    assert ws["B6"].value == "All" and ws["B7"].value == "Sales"
    assert ws.freeze_panes == "A12"
    rows = _data_rows(ws, 11)
    assert [r[0] for r in rows] == ["S-000", "S-003", "S-003", "S-003"]  # oldest day first
    assert rows[1][27] == "Es Kopi Calf Premium" and rows[2][27] == "Normal Sugar (PACKAGE)"
    assert rows[1][6] == datetime(2026, 9, 30)  # real Excel date
    assert ws.cell(row=13, column=7).number_format == "yyyy-mm-dd"
    assert rows[1][42] == 2500  # Bill Discount

    summary = list(wb["Ringkasan"].iter_rows(values_only=True))
    total = next(r for r in summary if r[0] == "TOTAL")
    # gross, -void, -other cost, -open, sales subtotal
    assert total[1:6] == (200000, -30000, -50000, 0, 120000)


def test_daily_report_export(client):
    job = _start(client, dateFrom="2026-09-29", dateTo="2026-09-30", report="daily")
    assert job["status"] == "done" and job["report"] == "daily"
    assert job["fileName"] == "Daily_Sales_Recapitulation_2026-09-29_to_2026-09-30.xlsx"
    wb = _workbook(client, job)
    assert wb.sheetnames == ["Report", "Ringkasan"]
    ws = wb["Report"]
    assert ws["A1"].value == "Daily Sales Recapitulation Report" and ws["B8"].value == "Daily"
    rows = list(ws.iter_rows(min_row=12, values_only=True))
    assert rows[0][:6] == ("Sales Date", "Sales Type", "Branch", "Number of Bill", "Pax Total", "Subtotal")
    assert [(r[0], r[2], r[3], r[5], r[16], r[18]) for r in rows[1:]] == [
        (datetime(2026, 9, 29), "Kopi Calf Supratman Bandung", 1, 20000, 20000, 20000),
        (datetime(2026, 9, 30), "Kopi Calf Supratman Bandung", 1, 100000, 90000, 90000),
    ]


def test_export_types_and_branch(client):
    job = _start(client, dateFrom="2026-09-30", dateTo="2026-09-30", type="other_cost")
    assert job["rows"] == 1 and job["fileName"].endswith("_other_cost.xlsx")
    rows = _data_rows(_workbook(client, job)["Report"], 11)
    assert rows[0][0] == "S-001" and rows[0][44] == "Kasir Pamulang"

    job = _start(client, dateFrom="2026-09-30", dateTo="2026-09-30", branch="TGP17")
    assert job["rows"] == 0  # TGP17 only has an other-cost sale


def test_export_splits_into_multiple_sheets(client, monkeypatch):
    monkeypatch.setattr(exports, "MAX_SHEET_ROWS", 12)  # 10 title rows -> 2 data rows on sheet 1
    job = _start(client, dateFrom="2026-09-29", dateTo="2026-09-30")
    assert job["sheets"] == 2
    wb = _workbook(client, job)
    assert wb.sheetnames == ["Report", "Report (2)", "Ringkasan"]
    assert len(_data_rows(wb["Report"], 11)) == 2
    assert len(_data_rows(wb["Report (2)"], 1)) == 2


def test_export_without_data_has_no_file(client):
    job = _start(client, dateFrom="2026-01-01", dateTo="2026-01-02")
    assert job["status"] == "done" and job["rows"] == 0
    assert job["fileName"] is None and job["downloadUrl"] is None
    assert client.get(f"/api/exports/{job['id']}/download").status_code == 409


def test_export_absolute_download_url(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "public_base_url", "https://portal-api.example.com/")
    job = _start(client, dateFrom="2026-09-30", dateTo="2026-09-30")
    assert job["downloadUrl"] == f"https://portal-api.example.com/api/exports/{job['id']}/download"


def test_export_default_range_when_no_dates(client):
    assert _start(client)["totalDays"] == get_settings().default_days + 1


def test_export_validation(client):
    assert client.post("/api/exports", json={"dateFrom": "30-09-2026", "dateTo": "2026-09-30"}).status_code == 422
    res = client.post("/api/exports", json={"dateFrom": "2026-10-01", "dateTo": "2026-09-30"})
    assert res.status_code == 422 and "error" in res.json()


def test_export_failure_is_reported(client, fake_db):
    fake_db.fail = True
    job = _start(client, dateFrom="2026-09-30", dateTo="2026-09-30")
    assert job["status"] == "error" and "database is down" in job["error"]
    assert not exports.job_file(job["id"]).exists()


def test_unknown_or_invalid_job_id(client):
    assert client.get("/api/exports/" + "0" * 32).status_code == 404
    assert client.get("/api/exports/../../etc/passwd").status_code == 404
    assert client.get("/api/exports/not-a-job/download").status_code == 404


def test_stale_running_job_reported_as_error(client):
    job_id = "a" * 32
    old = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
    state = {"id": job_id, "status": "running", "updatedAt": old, "fileName": "x.xlsx"}
    (exports._export_dir() / f"{job_id}.json").write_text(json.dumps(state), encoding="utf-8")
    assert client.get(f"/api/exports/{job_id}").json()["status"] == "error"


def _write_state(**state):
    job_id = state.setdefault("id", "b" * 32)
    now = datetime.now(timezone.utc).isoformat()
    state.setdefault("updatedAt", now)
    state.setdefault("createdAt", now)
    state.setdefault("fileName", "x.xlsx")
    (exports._export_dir() / f"{job_id}.json").write_text(json.dumps(state), encoding="utf-8")
    return job_id


def test_dead_export_process_reported_immediately(client, monkeypatch):
    monkeypatch.setattr(exports, "_process_alive", lambda pid: pid == 111)
    alive = _write_state(id="c" * 32, status="running", pid=111)
    dead = _write_state(id="d" * 32, status="running", pid=222)
    assert client.get(f"/api/exports/{alive}").json()["status"] == "running"
    body = client.get(f"/api/exports/{dead}").json()
    assert body["status"] == "error" and "berhenti" in body["error"]


def test_export_process_that_never_started(client):
    old = (datetime.now(timezone.utc) - timedelta(minutes=5)).isoformat()
    job_id = _write_state(status="queued", createdAt=old, updatedAt=old)
    assert client.get(f"/api/exports/{job_id}").json()["status"] == "error"


def test_run_job_records_its_pid(client, monkeypatch):
    seen = {}
    monkeypatch.setattr(exports, "_run", lambda job: seen.update(job))
    job_id = _write_state(status="queued", dateFrom="2026-09-30", dateTo="2026-09-30")
    exports.run_job(job_id)
    assert seen["pid"] == os.getpid()
    assert json.loads((exports._export_dir() / f"{job_id}.json").read_text())["pid"] == os.getpid()


def _google_configured(monkeypatch, uploads):
    s = get_settings()
    for key, value in (("google_drive_folder_id", "folder1"), ("google_oauth_client_id", "cid"),
                       ("google_oauth_client_secret", "secret"), ("google_oauth_refresh_token", "refresh")):
        monkeypatch.setattr(s, key, value)

    def upload(path, title, share_with, progress=lambda pct: None, folder=None):
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == ["Report", "Ringkasan"]
        progress(0.5)
        n = len(uploads) + 1
        uploads.append((title, share_with) if folder is None else (title, share_with, folder, len(list(wb["Report"].iter_rows(min_row=12)))))  # header on row 11
        return {"id": f"sheet{n}", "url": f"https://docs.google.com/spreadsheets/d/sheet{n}/edit", "sharedWith": share_with,
                "linkAccess": "view", "folderUrl": "https://drive.google.com/drive/folders/f1" if folder else None}

    monkeypatch.setattr(gsheets, "upload_as_sheet", upload)


def test_google_sheets_export_disabled_without_credentials(client):
    assert client.get("/api/exports").json()["googleSheets"] is False
    res = client.post("/api/exports", json={"dateFrom": "2026-09-30", "dateTo": "2026-09-30", "format": "gsheet"})
    assert res.status_code == 422 and "Google Sheets" in res.json()["error"]


def test_google_sheets_export(client, monkeypatch, fake_db):
    uploads = []
    _google_configured(monkeypatch, uploads)
    assert client.get("/api/exports").json()["googleSheets"] is True
    job = _start(client, dateFrom="2026-09-29", dateTo="2026-09-30", format="gsheet")
    assert job["status"] == "done" and job["format"] == "gsheet" and job["phase"] is None
    assert job["sheetUrl"] == "https://docs.google.com/spreadsheets/d/sheet1/edit"
    assert job["sheetSharedWith"] == "tester@kopicalf.co.id" and "ownerEmail" not in job
    assert job["sheetLinkAccess"] == "view" and len(job["sheetParts"]) == 1 and job["sheetFolderUrl"] is None
    assert uploads == [("Sales_Recapitulation_Detail_2026-09-29_to_2026-09-30", "tester@kopicalf.co.id")]
    assert job["downloadUrl"]  # the .xlsx stays downloadable as well
    done = next(a for a in fake_db.activity if a["action"] == "export.done")
    assert done["details"]["sheetUrl"] == job["sheetUrl"] and "Google Sheets" in done["summary"]


def test_google_sheets_export_in_parts(client, monkeypatch, fake_db):
    """Too many rows for one Google Sheet: parts cut at a day (or inside one), one Drive folder, shared by link."""
    uploads = []
    _google_configured(monkeypatch, uploads)
    monkeypatch.setattr(exports, "PART_TARGET_ROWS", 1)   # a new part after every day with rows
    monkeypatch.setattr(exports, "PART_MAX_ROWS", 2)      # and inside a day after 2 rows
    job = _start(client, dateFrom="2026-09-29", dateTo="2026-09-30", format="gsheet")
    assert job["status"] == "done" and job["rows"] == 4 and job["sheets"] == 3
    # 29 Sep: 1 row · 30 Sep: 3 rows -> cut after 2
    assert [(p["dateFrom"], p["dateTo"], p["rows"]) for p in job["sheetParts"]] == [
        ("2026-09-29", "2026-09-29", 1), ("2026-09-30", "2026-09-30", 2), ("2026-09-30", "2026-09-30", 1)]
    assert [u[2:] for u in uploads] == [(uploads[0][2], 1), (uploads[0][2], 2), (uploads[0][2], 1)]
    assert uploads[0][2].startswith("Sales_Recapitulation_Detail_2026-09-29_to_2026-09-30 (")   # one folder for all parts
    assert uploads[1][0].startswith("Sales_Recapitulation_Detail_2026-09-29_to_2026-09-30 - part 2 (30 Sep 2026")
    assert job["sheetUrl"] == job["sheetFolderUrl"] == "https://drive.google.com/drive/folders/f1"
    assert job["sheetLinkAccess"] == "view" and job["downloadUrl"] is None   # no single .xlsx of the whole range
    assert not list(exports._export_dir().glob(f"{job['id']}.part*.xlsx"))     # part files removed after upload
    done = next(a for a in fake_db.activity if a["action"] == "export.done")
    assert len(done["details"]["sheetParts"]) == 3


def test_google_sheets_upload_failure(client, monkeypatch):
    _google_configured(monkeypatch, [])

    def fail(*args, **kwargs):
        raise gsheets.SheetsError("Google Drive menolak upload (HTTP 403): quota")

    monkeypatch.setattr(gsheets, "upload_as_sheet", fail)
    job = _start(client, dateFrom="2026-09-30", dateTo="2026-09-30", format="gsheet")
    assert job["status"] == "error" and "HTTP 403" in job["error"] and job["sheetUrl"] is None


class _Resp:
    def __init__(self, body, status=200):
        self.body, self.status_code, self.text = body, status, str(body)

    def json(self):
        if isinstance(self.body, dict):
            return self.body
        raise ValueError("not json")


def test_google_sheets_via_apps_script(client, monkeypatch):
    import base64

    import requests

    s = get_settings()
    monkeypatch.setattr(s, "google_apps_script_url", "https://script.google.com/macros/s/x/exec")
    monkeypatch.setattr(s, "google_apps_script_key", "k" * 43)
    calls = []

    def post(url, json=None, timeout=None):
        calls.append(json)
        if json["action"] == "pair":
            return _Resp({"ok": True, "folderUrl": "https://drive.google.com/drive/folders/f1"})
        assert openpyxl.load_workbook(io.BytesIO(base64.b64decode(json["data"]))).sheetnames == ["Report", "Ringkasan"]
        return _Resp({"ok": True, "id": "s1", "url": "https://docs.google.com/spreadsheets/d/s1/edit", "sharedWith": json["shareWith"]})

    monkeypatch.setattr(requests, "post", post)
    assert gsheets.enabled() and gsheets.pair()["ok"]
    assert client.get("/api/exports").json()["googleSheets"] is True
    job = _start(client, dateFrom="2026-09-29", dateTo="2026-09-30", format="gsheet")
    assert job["status"] == "done" and job["sheetUrl"] == "https://docs.google.com/spreadsheets/d/s1/edit"
    upload = calls[-1]
    assert upload["key"] == "k" * 43 and upload["name"] == "Sales_Recapitulation_Detail_2026-09-29_to_2026-09-30"
    assert upload["shareWith"] == "tester@kopicalf.co.id" and upload["role"] == "writer"
    assert upload["linkAccess"] == "view" and upload["folder"] is None   # one part: no folder

    # refused by the script / not a JSON answer / file too large -> clear error on the job
    monkeypatch.setattr(requests, "post", lambda *a, **k: _Resp({"ok": False, "error": "unauthorized"}))
    assert "menolak: unauthorized" in _start(client, dateFrom="2026-09-30", dateTo="2026-09-30", format="gsheet")["error"]
    monkeypatch.setattr(requests, "post", lambda *a, **k: _Resp("<html>", 404))
    assert "bukan JSON" in _start(client, dateFrom="2026-09-30", dateTo="2026-09-30", format="gsheet")["error"]
    monkeypatch.setattr(gsheets, "APPS_SCRIPT_MAX_BYTES", 10)
    assert "melebihi batas Google Apps Script" in _start(client, dateFrom="2026-09-30", dateTo="2026-09-30", format="gsheet")["error"]


def test_maintenance_purges_activity_and_old_sheets(monkeypatch):
    import requests

    from app import database as db
    from app import maintenance

    s = get_settings()
    monkeypatch.setattr(s, "google_apps_script_url", "https://script.google.com/macros/s/x/exec")
    monkeypatch.setattr(s, "google_apps_script_key", "k" * 43)
    sql, sent = [], []

    class Conn:
        def execute(self, q, p=None):
            sql.append((q, p))
            return type("R", (), {"rowcount": 7})()

    class Tx:
        def __enter__(self):
            return Conn()

        def __exit__(self, *a):
            return False

    monkeypatch.setattr(db, "transaction", lambda timeout_ms=None: Tx())
    monkeypatch.setattr(requests, "post", lambda url, json=None, timeout=None: sent.append(json) or _Resp({"ok": True, "trashed": 3}))
    assert maintenance.purge_activity() == 7 and sql[0][1] == ("90 days",) and "activity_log" in sql[0][0]
    assert maintenance.purge_sheets() == 3 and sent[0]["action"] == "cleanup" and sent[0]["days"] == 30
    assert maintenance.main() == 0

    monkeypatch.setattr(s, "activity_retention_days", 0)
    monkeypatch.setattr(s, "gsheet_retention_days", 0)
    sql.clear(); sent.clear()
    assert maintenance.purge_activity() == 0 and maintenance.purge_sheets() == 0 and not sql and not sent

    # a failing clean-up is reported, the other still runs
    monkeypatch.setattr(s, "activity_retention_days", 90)
    monkeypatch.setattr(s, "gsheet_retention_days", 30)
    monkeypatch.setattr(requests, "post", lambda *a, **k: _Resp({"ok": False, "error": "unauthorized"}))
    assert maintenance.main() == 1 and sql
